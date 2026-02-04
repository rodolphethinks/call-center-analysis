"""Verify Korean character encoding in reports"""
import pandas as pd
from docx import Document
import sys

# Set UTF-8 output
sys.stdout.reconfigure(encoding='utf-8')

print("=" * 100)
print("VERIFYING KOREAN CHARACTER ENCODING")
print("=" * 100)

# Test Excel
print("\n📊 TESTING EXCEL FILE: report_unicode.xlsx")
print("-" * 100)
df = pd.read_excel('data/output/report_unicode.xlsx', sheet_name='Call Analysis')

# Find rows with Korean text
korean_rows = df[df['Agent Summary'].str.contains('인수증|안녕하세요|아|네', na=False)]

if len(korean_rows) > 0:
    print(f"\n✅ Found {len(korean_rows)} rows with Korean text")
    print("\nSample Korean text from first matching row:")
    row = korean_rows.iloc[0]
    print(f"  Call ID: {row['Call ID']}")
    print(f"  Agent Summary: {row['Agent Summary'][:150]}...")
    print(f"  Sentiment Examples: {row['Customer Sentiment Examples'][:150]}...")
else:
    print("❌ No Korean text found - this might indicate an encoding issue")

# Check for common corruption patterns
corrupted = df[df['Agent Summary'].str.contains('ㄱ|ㄴ|ㄷ|ㅁ|ㅂ|ㅅ|ㅇ', na=False, regex=True)]
if len(corrupted) > 0:
    print(f"\n⚠️  WARNING: Found {len(corrupted)} rows with potentially corrupted Korean text (isolated Jamo)")
    print("Sample corrupted text:")
    print(f"  {corrupted.iloc[0]['Agent Summary'][:100]}")
else:
    print("\n✅ No corrupted Korean text detected (no isolated Jamo characters)")

# Test Word
print("\n\n📄 TESTING WORD FILE: report_unicode.docx")
print("-" * 100)
doc = Document('data/output/report_unicode.docx')

korean_paras = [p for p in doc.paragraphs if any(char in p.text for char in '한글인수증안녕하세요아네')]

if korean_paras:
    print(f"\n✅ Found {len(korean_paras)} paragraphs with Korean text")
    print("\nSample Korean text:")
    for i, para in enumerate(korean_paras[:3]):
        print(f"  {i+1}. {para.text[:100]}...")
else:
    print("❌ No Korean text found in Word document")

# Check for corruption in Word
corrupted_paras = [p for p in doc.paragraphs if any(char in p.text for char in 'ㄱㄴㄷㅁㅂㅅㅇ') and '한글' not in p.text]
if corrupted_paras:
    print(f"\n⚠️  WARNING: Found {len(corrupted_paras)} paragraphs with potentially corrupted text")
else:
    print("\n✅ No corrupted text detected in Word document")

print("\n" + "=" * 100)
print("VERIFICATION COMPLETE")
print("=" * 100)
print("\n💡 If you see proper Korean text above (like 인수증, 안녕하세요), the encoding is working correctly!")
print("   Open the Excel and Word files to verify they display properly in the applications.")
