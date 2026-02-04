"""Test script to verify report encoding"""
import pandas as pd
from docx import Document

# Test Excel
print("=" * 80)
print("TESTING EXCEL FILE")
print("=" * 80)
df = pd.read_excel('data/output/report_final.xlsx', sheet_name='Call Analysis')
print(f"\nTotal rows: {len(df)}")
print(f"\nColumns: {list(df.columns)}")
print(f"\nFirst 3 rows with Korean text:")
for idx in range(min(3, len(df))):
    row = df.iloc[idx]
    print(f"\n--- Row {idx + 1} ---")
    print(f"Call ID: {row['Call ID']}")
    print(f"Duration: {row['Duration']:.2f} seconds ({row['Duration']/60:.1f} minutes)")
    print(f"Customer Summary: {row['Customer Summary'][:200]}...")
    print(f"Sentiment: {row['Customer Sentiment Trajectory']}")

# Test Word
print("\n" + "=" * 80)
print("TESTING WORD FILE")
print("=" * 80)
doc = Document('data/output/report_final.docx')
print(f"\nTotal paragraphs: {len(doc.paragraphs)}")
print(f"\nFirst 10 paragraphs:")
for i, para in enumerate(doc.paragraphs[:10]):
    if para.text.strip():
        print(f"{i}: {para.text[:100]}")

print("\n✅ Test complete! Check if Korean characters display correctly above.")
